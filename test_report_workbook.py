"""Tests for the accountant-ready reconciliation workbook.

Synthetic data only. The workbook is generated from the real SQLite
database, reopened, and checked against the persisted run.
"""

from __future__ import annotations

import io
import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent))

from agent import store  # noqa: E402

import db  # noqa: E402
import demo_workflow  # noqa: E402
import milaan_web  # noqa: E402
import report_workbook  # noqa: E402

_CHARSET = "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ"


def _checksum(first_14: str) -> str:
    total = 0
    for position, character in enumerate(first_14):
        value = _CHARSET.index(character.upper())
        product = value * (1 if position % 2 == 0 else 2)
        total += product // 36 + product % 36
    return _CHARSET[(36 - total % 36) % 36]


MALICIOUS_GSTIN = "27AAACE3456D1Z" + _checksum("27AAACE3456D1Z")


class ReportWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="milaan-report-"))
        self.old_db = os.environ.get("DATABASE_PATH")
        self.old_agent = os.environ.get("AGENT_DB_PATH")
        os.environ["DATABASE_PATH"] = str(self.tmp / "milaan.db")
        os.environ["AGENT_DB_PATH"] = str(self.tmp / "agent.db")
        store.reset_initialized()
        db.init_db()
        self.case = demo_workflow.load_demo_workspace()
        self.content = report_workbook.build_reconciliation_workbook(
            self.case["run_id"]
        )
        self.app = milaan_web.create_app()
        self.client = self.app.test_client()

    def tearDown(self):
        if self.old_db is None:
            os.environ.pop("DATABASE_PATH", None)
        else:
            os.environ["DATABASE_PATH"] = self.old_db
        if self.old_agent is None:
            os.environ.pop("AGENT_DB_PATH", None)
        else:
            os.environ["AGENT_DB_PATH"] = self.old_agent
        store.reset_initialized()
        shutil.rmtree(self.tmp, ignore_errors=True)

    def workbook(self):
        return load_workbook(io.BytesIO(self.content), data_only=True)

    def test_required_sheets_and_self_validation(self):
        workbook = self.workbook()
        self.assertEqual(
            workbook.sheetnames, list(report_workbook.REQUIRED_SHEETS)
        )
        self.assertEqual(
            report_workbook.validate_workbook(
                self.content, self.case["run_id"]
            ),
            [],
        )
        workbook.close()

    def test_counts_and_metadata_match_persisted_run(self):
        run = db.get_reconciliation_run(self.case["run_id"])
        workbook = self.workbook()
        counts = {
            name: max(workbook[name].max_row - 1, 0)
            for name in report_workbook.REQUIRED_SHEETS[1:-1]
        }
        self.assertEqual(counts["Matched"], run["matched_count"])
        self.assertEqual(
            counts["Purchase Missing from 2B"], run["missing_count"]
        )
        self.assertEqual(
            counts["2B Missing from Books"],
            run["in_2b_not_in_books_count"],
        )
        self.assertEqual(
            counts["Duplicate Near-match Review"], run["needs_review_count"]
        )

        summary = {}
        for row in workbook["Summary"].iter_rows(values_only=True):
            if row and len(row) >= 2 and isinstance(row[0], str):
                summary.setdefault(row[0], row[1])
        self.assertEqual(summary["Run ID"], self.case["run_id"])
        self.assertEqual(summary["GSTR-2B batch"], self.case["batch_id"])
        self.assertEqual(summary["Matched"], run["matched_count"])
        self.assertEqual(
            summary["Purchase missing from 2B"], run["missing_count"]
        )
        self.assertEqual(
            summary["2B rows missing from books"],
            run["in_2b_not_in_books_count"],
        )
        workbook.close()

    def test_formula_injection_is_escaped_everywhere(self):
        with db.db() as conn:
            supplier_id = conn.execute(
                """INSERT INTO suppliers
                       (supplier_name, gstin, contact_phone)
                   VALUES (?, ?, ?)""",
                ("=HYPERLINK(\"http://example.test\",\"x\")",
                 MALICIOUS_GSTIN, "9999999999"),
            ).lastrowid
            conn.execute(
                """INSERT INTO invoices
                       (supplier_id, invoice_number, invoice_date, tax_period,
                        taxable_value, gst_rate, cgst, sgst, igst, total_gst,
                        invoice_total, validation_status)
                   VALUES (?, ?, '2026-03-05', '2026-03', 100000, 18,
                           9000, 9000, 0, 18000, 118000, 'valid')""",
                (supplier_id, "  =1+1"),
            )
        summary = db.reconcile_gstr2b_import(self.case["batch_id"], 0)
        content = report_workbook.build_reconciliation_workbook(
            summary["run_id"]
        )
        self.assertEqual(
            report_workbook.validate_workbook(
                content, summary["run_id"]
            ),
            [],
        )
        workbook = load_workbook(io.BytesIO(content), data_only=True)
        escaped_invoice = False
        escaped_supplier = False
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows(values_only=True):
                for value in row:
                    if not isinstance(value, str) or not value:
                        continue
                    self.assertFalse(
                        value.startswith(("=", "+", "@")),
                        f"formula-like value in {worksheet.title}: {value!r}",
                    )
                    if value == "'  =1+1":
                        escaped_invoice = True
                    if value.startswith("'=HYPERLINK"):
                        escaped_supplier = True
        self.assertTrue(escaped_invoice)
        self.assertTrue(escaped_supplier)
        workbook.close()

    def test_validator_detects_count_and_formula_tampering(self):
        workbook = load_workbook(io.BytesIO(self.content))
        matched = workbook["Matched"]
        if matched.max_row > 1:
            matched.delete_rows(2)
        workbook["Audit Metadata"]["B2"] = "=1+1"
        buffer = io.BytesIO()
        workbook.save(buffer)
        workbook.close()
        problems = report_workbook.validate_workbook(
            buffer.getvalue(), self.case["run_id"]
        )
        self.assertTrue(
            any("rows" in problem for problem in problems), problems
        )
        self.assertTrue(
            any("formula" in problem for problem in problems), problems
        )

    def test_export_endpoints_return_valid_workbook(self):
        response = self.client.get(
            f"/api/reconciliation/{self.case['run_id']}/workbook.xlsx"
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.mimetype,
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet",
        )
        self.assertTrue(response.data.startswith(b"PK"))
        workbook = load_workbook(io.BytesIO(response.data), data_only=True)
        self.assertIn("Audit Metadata", workbook.sheetnames)
        workbook.close()

        latest = self.client.get("/api/export/reconciliation.xlsx")
        self.assertEqual(latest.status_code, 200)
        self.assertTrue(latest.data.startswith(b"PK"))

        missing = self.client.get(
            "/api/reconciliation/9999/workbook.xlsx"
        )
        self.assertEqual(missing.status_code, 404)


if __name__ == "__main__":
    unittest.main()
