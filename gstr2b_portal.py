"""Dedicated adapter for GST-portal GSTR-2B B2B Excel exports.

This is separate from the documented canonical CSV/XLSX schema in
``gstr2b.py``. Portal workbooks typically have title rows, a B2B sheet,
Indian-formatted amounts, and extra ITC-evidence columns. Ambiguous or
unsupported layouts fail with an actionable error instead of guessing.
"""

from __future__ import annotations

import io
import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

import gstr2b
import validation

SUPPORTED_SHEET_NAMES = frozenset({
    "b2b",
    "b2b invoices",
    "b2b invoice",
    "b2ba",
})
UNSUPPORTED_BUT_KNOWN_SHEETS = frozenset({
    "b2b-cdnr",
    "b2b cdnr",
    "cdnr",
    "cdnra",
    "b2b-cdnra",
    "isd",
    "isda",
    "impg",
    "impgsez",
    "itc rev",
    "itcrev",
    "hsn",
    "docs",
})

# Explicit portal header labels -> canonical field. Ambiguous labels are
# omitted on purpose so they fail rather than silently mapping.
PORTAL_HEADERS = {
    "gstin_of_supplier": "supplier_gstin",
    "gstin_of_the_supplier": "supplier_gstin",
    "supplier_gstin": "supplier_gstin",
    "invoice_number": "invoice_number",
    "invoice_no": "invoice_number",
    "invoice_num": "invoice_number",
    "document_number": "invoice_number",
    "invoice_date": "invoice_date",
    "invoice_dt": "invoice_date",
    "document_date": "invoice_date",
    "taxable_value": "taxable_value",
    "taxable_value_₹": "taxable_value",
    "taxable_value_(₹)": "taxable_value",
    "taxable_amount": "taxable_value",
    "integrated_tax": "igst",
    "integrated_tax_₹": "igst",
    "integrated_tax_(₹)": "igst",
    "igst": "igst",
    "igst_amount": "igst",
    "central_tax": "cgst",
    "central_tax_₹": "cgst",
    "central_tax_(₹)": "cgst",
    "cgst": "cgst",
    "cgst_amount": "cgst",
    "state/ut_tax": "sgst",
    "state_ut_tax": "sgst",
    "state/ut_tax_₹": "sgst",
    "state/ut_tax_(₹)": "sgst",
    "sgst": "sgst",
    "sgst_amount": "sgst",
    "cess": "cess",
    "cess_₹": "cess",
    "cess_(₹)": "cess",
    "gstr_1/iff/gstr_5_period": "tax_period",
    "gstr_1_iff_gstr_5_period": "tax_period",
    "return_period": "tax_period",
    "tax_period": "tax_period",
    "itc_availability": "itc_availability",
    "itc_available": "itc_availability",
    "reason": "itc_reason",
    "reason_for_itc_unavailability": "itc_reason",
    "supply_attract_reverse_charge": "reverse_charge",
    "reverse_charge": "reverse_charge",
    "invoice_type": "document_type",
    "document_type": "document_type",
    "irn": "irn",
    "invoice_reference_number": "irn",
    "invoice_value": "invoice_value",
    "invoice_value_₹": "invoice_value",
    "invoice_value_(₹)": "invoice_value",
}

REQUIRED_PORTAL_FIELDS = (
    "supplier_gstin",
    "invoice_number",
    "invoice_date",
    "taxable_value",
    "cgst",
    "sgst",
    "igst",
)

OPTIONAL_PORTAL_FIELDS = (
    "tax_period",
    "cess",
    "itc_availability",
    "itc_reason",
    "reverse_charge",
    "document_type",
    "irn",
    "invoice_value",
    "total_gst",
)

_MONTH_NAMES = {
    "jan": 1, "january": 1,
    "feb": 2, "february": 2,
    "mar": 3, "march": 3,
    "apr": 4, "april": 4,
    "may": 5,
    "jun": 6, "june": 6,
    "jul": 7, "july": 7,
    "aug": 8, "august": 8,
    "sep": 9, "sept": 9, "september": 9,
    "oct": 10, "october": 10,
    "nov": 11, "november": 11,
    "dec": 12, "december": 12,
}


def parse_portal_money_to_paise(value: Any) -> int:
    """Portal amounts may include rupee signs and grouping commas."""
    if isinstance(value, str):
        text = (
            value.strip()
            .replace("₹", "")
            .replace("Rs.", "")
            .replace("Rs", "")
            .replace(",", "")
            .replace("\u00a0", "")
        )
        if text.startswith("(") and text.endswith(")"):
            raise gstr2b.MoneyParseError(
                "negative monetary value", code="NEGATIVE_AMOUNT"
            )
        return gstr2b.parse_money_to_paise(text)
    return gstr2b.parse_money_to_paise(value)


def parse_tax_period(value: Any) -> str | None:
    """Accept YYYY-MM, YYYY/MM, MMYYYY, YYYYMM, and Mon-YYYY forms."""
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return f"{value.year:04d}-{value.month:02d}"
    text = gstr2b._cell_text(value)
    if not text:
        return None
    text = text.replace("\u00a0", " ").strip()
    match = re.fullmatch(r"(\d{4})[-/](0[1-9]|1[0-2])", text)
    if match:
        return f"{match.group(1)}-{match.group(2)}"
    match = re.fullmatch(r"(0[1-9]|1[0-2])[-/](\d{4})", text)
    if match:
        return f"{match.group(2)}-{match.group(1)}"
    match = re.fullmatch(r"(0[1-9]|1[0-2])(\d{4})", text)
    if match:
        return f"{match.group(2)}-{match.group(1)}"
    match = re.fullmatch(r"(\d{4})(0[1-9]|1[0-2])", text)
    if match:
        return f"{match.group(1)}-{match.group(2)}"
    match = re.fullmatch(r"([A-Za-z]+)\s*[-/]\s*(\d{4})", text)
    if match:
        month = _MONTH_NAMES.get(match.group(1).lower())
        if month:
            return f"{match.group(2)}-{month:02d}"
    return None


def parse_invoice_date(value: Any) -> str | None:
    if value is None or gstr2b._cell_text(value) == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = gstr2b._cell_text(value)
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        try:
            return date.fromisoformat(text).isoformat()
        except ValueError:
            return None
    match = re.fullmatch(r"(\d{1,2})[-/](\d{1,2})[-/](\d{4})", text)
    if match:
        day, month, year = (int(match.group(1)), int(match.group(2)), int(match.group(3)))
        try:
            return date(year, month, day).isoformat()
        except ValueError:
            return None
    return None


def _normalize_portal_header(value: Any) -> str:
    text = gstr2b._normalize_header(value)
    text = text.replace("₹", "").replace("(", "").replace(")", "")
    text = re.sub(r"_+", "_", text).strip("_")
    return text


def _yes_no(value: Any) -> str | None:
    text = gstr2b._cell_text(value).upper()
    if not text:
        return None
    if text in {"Y", "YES", "TRUE", "AVAILABLE", "1"}:
        return "Y"
    if text in {"N", "NO", "FALSE", "UNAVAILABLE", "0"}:
        return "N"
    return text


def looks_like_portal_headers(headers: list[Any]) -> bool:
    names = {_normalize_portal_header(cell) for cell in headers if gstr2b._cell_text(cell)}
    mapped = {PORTAL_HEADERS[name] for name in names if name in PORTAL_HEADERS}
    return "supplier_gstin" in mapped and "invoice_number" in mapped and "taxable_value" in mapped


PORTAL_ONLY_HEADERS = frozenset(
    set(PORTAL_HEADERS) - set(gstr2b.COLUMN_ALIASES)
)


def looks_like_portal_csv_table(rows: list[list[Any]]) -> bool:
    """Return True for a portal CSV header, not a broken canonical template.

    A canonical template with one missing column must stay a fatal canonical
    error. Portal CSVs are recognised by at least one portal-specific label
    such as ``Integrated Tax`` or ``GSTR-1/IFF/GSTR-5 Period``.
    """
    for row in rows[:30]:
        if not looks_like_portal_headers(row):
            continue
        names = {
            _normalize_portal_header(cell)
            for cell in row
            if gstr2b._cell_text(cell)
        }
        if names & PORTAL_ONLY_HEADERS:
            return True
    return False


def _sheet_kind(name: str) -> str:
    key = re.sub(r"[\s_]+", " ", (name or "").strip().lower())
    if key in SUPPORTED_SHEET_NAMES or key.startswith("b2b") and "cdn" not in key and "isd" not in key:
        if "cdn" in key:
            return "unsupported"
        return "supported"
    if key in UNSUPPORTED_BUT_KNOWN_SHEETS or "cdn" in key or key.startswith("isd") or key.startswith("impg"):
        return "unsupported"
    return "other"


def _header_index(headers: list[Any]) -> dict[str, int]:
    index: dict[str, int] = {}
    collisions: list[str] = []
    unknown: list[str] = []
    for position, header in enumerate(headers):
        original = _normalize_portal_header(header)
        if not original:
            continue
        field = PORTAL_HEADERS.get(original)
        if field is None:
            unknown.append(gstr2b._cell_text(header) or original)
            continue
        if field in index:
            collisions.append(field)
            continue
        index[field] = position
    if collisions:
        raise ValueError(
            "ambiguous GSTR-2B columns for: "
            + ", ".join(sorted(set(collisions)))
            + ". The file was not imported."
        )
    missing = [name for name in REQUIRED_PORTAL_FIELDS if name not in index]
    if missing:
        raise ValueError(
            "portal GSTR-2B B2B sheet is missing required column(s): "
            + ", ".join(missing)
        )
    return index


def _issue(code: str, field: str, message: str) -> validation.Issue:
    return validation.Issue(code=code, field=field, message=message)


def _parse_portal_row(
    *,
    row_number: int,
    row: list[Any],
    headers: list[str],
    column_index: dict[str, int],
    seen_keys: set[tuple[str, str, str]],
    source_sheet: str,
    source_row: int,
    fallback_period: str | None,
) -> gstr2b.ParsedRow:
    errors: list[validation.Issue] = []
    warnings: list[validation.Issue] = []

    def cell(field: str) -> Any:
        position = column_index.get(field)
        if position is None or position >= len(row):
            return None
        return row[position]

    supplier_gstin_text = gstr2b._cell_text(cell("supplier_gstin"))
    supplier_gstin: str | None = None
    if not supplier_gstin_text:
        errors.append(_issue("REQUIRED_FIELD", "supplier_gstin", "supplier_gstin is required."))
    else:
        gstin_check = validation.validate_gstin(supplier_gstin_text)
        supplier_gstin = gstin_check.normalized or supplier_gstin_text
        errors.extend(gstin_check.errors)

    invoice_number = gstr2b._cell_text(cell("invoice_number")) or None
    if invoice_number is None:
        errors.append(_issue("REQUIRED_FIELD", "invoice_number", "invoice_number is required."))

    raw_date = cell("invoice_date")
    invoice_date = parse_invoice_date(raw_date)
    if gstr2b._cell_text(raw_date) == "":
        errors.append(_issue("REQUIRED_FIELD", "invoice_date", "invoice_date is required."))
    elif invoice_date is None:
        errors.append(_issue(
            "INVALID_DATE", "invoice_date",
            "invoice_date must be a real calendar date (YYYY-MM-DD or DD-MM-YYYY).",
        ))

    raw_period = cell("tax_period") if "tax_period" in column_index else None
    tax_period = parse_tax_period(raw_period) if raw_period is not None else None
    if tax_period is None and fallback_period:
        tax_period = fallback_period
        warnings.append(_issue(
            "TAX_PERIOD_FROM_FILE", "tax_period",
            f"tax_period taken from workbook metadata ({fallback_period}).",
        ))
    if tax_period is None and invoice_date:
        tax_period = invoice_date[:7]
        warnings.append(_issue(
            "TAX_PERIOD_FROM_INVOICE_DATE", "tax_period",
            "tax_period was derived from invoice_date; human review is required.",
        ))
    if tax_period is None:
        errors.append(_issue(
            "REQUIRED_FIELD", "tax_period",
            "tax_period is required (YYYY-MM, MMYYYY, or Mon-YYYY).",
        ))

    if invoice_date and tax_period and invoice_date[:7] != tax_period:
        warnings.append(_issue(
            "DATE_PERIOD_INCONSISTENT", "invoice_date",
            f"invoice_date {invoice_date} is not in tax period {tax_period}.",
        ))

    money_fields = ("taxable_value", "cgst", "sgst", "igst", "cess")
    money: dict[str, int | None] = {field: None for field in money_fields}
    for field in money_fields:
        raw_value = cell(field) if field in column_index or field != "cess" else None
        if field == "cess" and "cess" not in column_index:
            money[field] = 0
            continue
        if raw_value is None or gstr2b._cell_text(raw_value) == "":
            if field == "cess":
                money[field] = 0
                continue
            errors.append(_issue("REQUIRED_FIELD", field, f"{field} is required."))
            continue
        try:
            money[field] = parse_portal_money_to_paise(raw_value)
        except gstr2b.MoneyParseError as exc:
            errors.append(_issue(exc.code, field, str(exc)))

    total_gst: int | None = None
    if all(money[field] is not None for field in ("cgst", "sgst", "igst")):
        total_gst = money["cgst"] + money["sgst"] + money["igst"]
        if money["cess"]:
            total_gst += money["cess"]
        if "total_gst" in column_index:
            raw_total = cell("total_gst")
            if raw_total is not None and gstr2b._cell_text(raw_total):
                try:
                    stated = parse_portal_money_to_paise(raw_total)
                    if stated != total_gst:
                        errors.append(_issue(
                            "TAX_TOTAL_MISMATCH", "total_gst",
                            f"total_gst ({stated} paise) must equal "
                            f"cgst+sgst+igst+cess ({total_gst} paise).",
                        ))
                except gstr2b.MoneyParseError as exc:
                    errors.append(_issue(exc.code, "total_gst", str(exc)))
        has_intra = (money["cgst"] or 0) > 0 or (money["sgst"] or 0) > 0
        has_inter = (money["igst"] or 0) > 0
        if has_intra and has_inter:
            errors.append(_issue(
                "TAX_STRUCTURE_CONFLICT", "igst",
                "CGST/SGST and IGST cannot both be charged on the same taxable amount.",
            ))
        elif has_intra and money["cgst"] != money["sgst"]:
            errors.append(_issue(
                "CGST_SGST_MISMATCH", "cgst",
                f"CGST ({money['cgst']} paise) and SGST ({money['sgst']} paise) must be equal.",
            ))

    itc_availability = _yes_no(cell("itc_availability")) if "itc_availability" in column_index else None
    itc_reason = gstr2b._cell_text(cell("itc_reason")) or None if "itc_reason" in column_index else None
    reverse_charge = _yes_no(cell("reverse_charge")) if "reverse_charge" in column_index else None
    document_type = gstr2b._cell_text(cell("document_type")) or None if "document_type" in column_index else None
    irn = gstr2b._cell_text(cell("irn")) or None if "irn" in column_index else None

    if supplier_gstin and invoice_number and tax_period:
        key = (supplier_gstin, invoice_number, tax_period)
        if key in seen_keys:
            errors.append(_issue(
                "DUPLICATE_ROW", "invoice_number",
                "Another row in this file has the same supplier GSTIN, invoice number and tax period.",
            ))
        else:
            seen_keys.add(key)

    if warnings and not errors:
        status = validation.STATUS_NEEDS_REVIEW if any(
            issue.code in {"TAX_PERIOD_FROM_INVOICE_DATE", "DATE_PERIOD_INCONSISTENT"}
            for issue in warnings
        ) else validation.STATUS_VALID
        if status == validation.STATUS_NEEDS_REVIEW:
            # Staged rows only allow valid/invalid/needs_review; keep review.
            pass
    else:
        status = validation.STATUS_INVALID if errors else validation.STATUS_VALID

    extra = {
        "source_sheet": source_sheet,
        "source_row": source_row,
        "cess": money["cess"],
        "itc_availability": itc_availability,
        "itc_reason": itc_reason,
        "reverse_charge": reverse_charge,
        "document_type": document_type,
        "irn": irn,
    }
    raw = json.loads(gstr2b._raw_data_json(headers, row))
    raw.update({key: value for key, value in extra.items() if value is not None})

    return gstr2b.ParsedRow(
        row_number=row_number,
        supplier_gstin=supplier_gstin,
        invoice_number=invoice_number,
        invoice_date=invoice_date,
        tax_period=tax_period,
        taxable_value=money["taxable_value"],
        cgst=money["cgst"],
        sgst=money["sgst"],
        igst=money["igst"],
        total_gst=total_gst,
        validation_status=status,
        errors=tuple(errors),
        warnings=tuple(warnings),
        raw_data=json.dumps(raw, ensure_ascii=False, default=str),
        cess=money["cess"],
        itc_availability=itc_availability,
        itc_reason=itc_reason,
        reverse_charge=reverse_charge,
        document_type=document_type,
        irn=irn,
        source_sheet=source_sheet,
        source_row=source_row,
    )


def try_parse_portal_workbook(content: bytes, file_name: str) -> gstr2b.ParsedGstr2b | None:
    """Return a ParsedGstr2b when the workbook is a supported portal export.

    Returns None when the file is not portal-shaped, so the canonical parser
    can still handle simple one-sheet templates.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ValueError("XLSX support is unavailable (openpyxl is not installed)") from exc
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        return None

    try:
        sheet_notes: list[str] = []
        supported: list[tuple[str, list[list[Any]]]] = []
        unsupported_present: list[str] = []
        for sheet in workbook.worksheets:
            kind = _sheet_kind(sheet.title)
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            while rows and gstr2b._is_blank_row(rows[-1]):
                rows.pop()
            if kind == "unsupported":
                if any(not gstr2b._is_blank_row(row) for row in rows[1:] or rows):
                    unsupported_present.append(sheet.title)
                continue
            header_at = None
            for index, row in enumerate(rows[:30]):
                if looks_like_portal_headers(row):
                    header_at = index
                    break
            if header_at is None:
                if kind == "supported":
                    sheet_notes.append(sheet.title)
                continue
            supported.append((sheet.title, rows))
        if not supported:
            return None
        return _parse_supported_tables(
            supported,
            file_name=file_name,
            content=content,
            source_format="xlsx",
            skipped_sheets=unsupported_present,
        )
    finally:
        workbook.close()


def parse_portal_csv_rows(
    rows: list[list[Any]], content: bytes, file_name: str
) -> gstr2b.ParsedGstr2b | None:
    """Parse a portal-shaped CSV table (title rows are allowed)."""
    return _parse_supported_tables(
        [("CSV", rows)],
        file_name=file_name,
        content=content,
        source_format="csv",
        skipped_sheets=[],
    )


def _parse_supported_tables(
    tables: list[tuple[str, list[list[Any]]]],
    *,
    file_name: str,
    content: bytes,
    source_format: str,
    skipped_sheets: list[str],
    fallback_period: str | None = None,
) -> gstr2b.ParsedGstr2b | None:
    located: list[tuple[str, list[list[Any]], int]] = []
    for sheet_name, rows in tables:
        header_at = None
        for index, row in enumerate(rows[:30]):
            if looks_like_portal_headers(row):
                header_at = index
                break
        if header_at is None:
            continue
        located.append((sheet_name, rows, header_at))
    if not located:
        return None

    parsed_rows: list[gstr2b.ParsedRow] = []
    seen_keys: set[tuple[str, str, str]] = set()
    unknown_headers: list[str] = []
    row_number = 0
    for sheet_name, rows, header_at in located:
        headers = [gstr2b._cell_text(cell) for cell in rows[header_at]]
        for header in headers:
            key = _normalize_portal_header(header)
            if key and key not in PORTAL_HEADERS and header:
                unknown_headers.append(header)
        column_index = _header_index(rows[header_at])
        for offset, row in enumerate(rows[header_at + 1:], start=header_at + 2):
            if gstr2b._is_blank_row(row):
                continue
            row_number += 1
            parsed_rows.append(
                _parse_portal_row(
                    row_number=row_number,
                    row=row,
                    headers=headers,
                    column_index=column_index,
                    seen_keys=seen_keys,
                    source_sheet=sheet_name,
                    source_row=offset,
                    fallback_period=fallback_period,
                )
            )
    if not parsed_rows:
        raise ValueError(
            "portal GSTR-2B workbook has a B2B header but no data rows"
        )
    unique_unknown = tuple(sorted(set(unknown_headers)))
    return gstr2b.ParsedGstr2b(
        file_name=file_name,
        source_format=source_format,
        file_hash=gstr2b.sha256_hex(content),
        file_size=len(content),
        rows=tuple(parsed_rows),
        parser_kind="portal_b2b",
        unsupported_fields=unique_unknown,
        sheets_used=tuple(name for name, _rows, _index in located),
        skipped_sheets=tuple(skipped_sheets),
    )


# --- GST-portal GSTR-2B JSON ------------------------------------------------
# The portal JSON nests invoices as data.docdata.b2b[].inv[].items[] (older
# and GSTR-2A-shaped exports use data.b2b / b2b and itms[].itm_det). Each
# invoice is flattened into one portal-shaped row and validated by the same
# _parse_portal_row rules as the Excel export, so amounts, ITC availability,
# reverse charge and cess are handled identically. Amounts are summed with
# Decimal; a missing taxable value is never replaced by the invoice value.

JSON_ROW_HEADERS = [
    "GSTIN of supplier", "Invoice number", "Invoice Date", "Taxable Value",
    "Integrated Tax", "Central Tax", "State/UT Tax", "Cess",
    "GSTR-1/IFF/GSTR-5 Period", "ITC Availability", "Reason",
    "Supply Attract Reverse Charge", "Invoice Value", "IRN", "Invoice type",
]
_JSON_ITEM_KEYS = {
    "taxable_value": ("txval",),
    "igst": ("igst", "iamt"),
    "cgst": ("cgst", "camt"),
    "sgst": ("sgst", "samt"),
    "cess": ("cess", "csamt"),
}
_JSON_SECTIONS_SKIPPED = ("b2ba", "cdnr", "cdnra", "isd", "isda", "impg", "impgsez")


def _json_b2b_container(data: Any) -> tuple[Any, dict] | None:
    """Return (b2b list, container dict) for a portal-shaped JSON document."""
    if not isinstance(data, dict):
        return None
    candidates = []
    inner = data.get("data")
    if isinstance(inner, dict):
        if isinstance(inner.get("docdata"), dict):
            candidates.append(inner["docdata"])
        candidates.append(inner)
    if isinstance(data.get("docdata"), dict):
        candidates.append(data["docdata"])
    candidates.append(data)
    for container in candidates:
        if "b2b" in container:
            return container["b2b"], container
    return None


def _sum_json_amount(items: list[dict], field: str, path: str) -> str:
    """Sum one tax field over invoice items as an exact rupee string.

    Absent keys count as zero (the portal omits zero taxes). Anything that is
    not a plain number is returned verbatim so row validation rejects it.
    """
    total = Decimal("0")
    found = False
    for item in items:
        for key in _JSON_ITEM_KEYS[field]:
            if key not in item:
                continue
            value = item[key]
            found = True
            if isinstance(value, bool) or not isinstance(value, (int, float, str)):
                return f"invalid value at {path}.{key}"
            try:
                amount = Decimal(str(value).strip())
            except InvalidOperation:
                return str(value)
            if not amount.is_finite():
                return str(value)
            total += amount
            break
    if field == "taxable_value" and not found:
        return ""
    return format(total, "f")


def try_parse_portal_json(content: bytes, file_name: str) -> gstr2b.ParsedGstr2b | None:
    """Parse GST-portal GSTR-2B JSON; return None when the JSON is not portal-shaped."""
    try:
        data = json.loads(content.decode("utf-8-sig"))
    except (UnicodeDecodeError, ValueError) as exc:
        raise ValueError(f"JSON file is not valid JSON text: {exc}") from exc
    located = _json_b2b_container(data)
    if located is None:
        return None
    b2b, container = located
    if not isinstance(b2b, list):
        raise ValueError("GSTR-2B JSON: 'b2b' must be a list of suppliers. Nothing was imported.")

    root = data.get("data") if isinstance(data.get("data"), dict) else data
    fallback_period = parse_tax_period(root.get("rtnprd") or data.get("rtnprd"))
    skipped = [
        key for key in _JSON_SECTIONS_SKIPPED
        if isinstance(container.get(key), list) and container.get(key)
    ]

    rows: list[list[Any]] = []
    for s_index, supplier in enumerate(b2b):
        s_path = f"b2b[{s_index}]"
        if not isinstance(supplier, dict):
            raise ValueError(f"GSTR-2B JSON: {s_path} must be an object. Nothing was imported.")
        invoices = supplier.get("inv", [])
        if not isinstance(invoices, list):
            raise ValueError(f"GSTR-2B JSON: {s_path}.inv must be a list. Nothing was imported.")
        for i_index, inv in enumerate(invoices):
            i_path = f"{s_path}.inv[{i_index}]"
            if not isinstance(inv, dict):
                raise ValueError(f"GSTR-2B JSON: {i_path} must be an object. Nothing was imported.")
            raw_items = inv.get("items", inv.get("itms", []))
            if not isinstance(raw_items, list):
                raise ValueError(f"GSTR-2B JSON: {i_path}.items must be a list. Nothing was imported.")
            items = []
            for t_index, item in enumerate(raw_items):
                if not isinstance(item, dict):
                    raise ValueError(
                        f"GSTR-2B JSON: {i_path}.items[{t_index}] must be an object. Nothing was imported."
                    )
                detail = item.get("itm_det", item)
                if not isinstance(detail, dict):
                    raise ValueError(
                        f"GSTR-2B JSON: {i_path}.items[{t_index}].itm_det must be an object. "
                        "Nothing was imported."
                    )
                items.append(detail)
            amounts = {field: _sum_json_amount(items, field, f"{i_path}.items") for field in _JSON_ITEM_KEYS}
            rows.append([
                supplier.get("ctin") or "",
                inv.get("inum") if inv.get("inum") is not None else "",
                inv.get("dt") or inv.get("idt") or "",
                amounts["taxable_value"],
                amounts["igst"],
                amounts["cgst"],
                amounts["sgst"],
                amounts["cess"],
                supplier.get("supprd") or "",
                inv.get("itcavl") or "",
                inv.get("rsn") or "",
                inv.get("rev") or "",
                inv.get("val") if inv.get("val") is not None else "",
                inv.get("irn") or "",
                inv.get("typ") or inv.get("inv_typ") or "",
            ])
    if not rows:
        raise ValueError("GSTR-2B JSON contains no B2B invoices. Nothing was imported.")
    return _parse_supported_tables(
        [("JSON b2b", [JSON_ROW_HEADERS, *rows])],
        file_name=file_name,
        content=content,
        source_format="json",
        skipped_sheets=skipped,
        fallback_period=fallback_period,
    )
