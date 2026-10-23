"""Accountant-ready XLSX workbook for one persisted reconciliation run.

The workbook is a read-only projection of stored evidence. It never changes
invoices, GSTR-2B rows, reconciliation results, follow-ups or delivery state.

All monetary amounts are derived from integer paise in SQLite. Display
columns are rupees with two decimals; difference columns stay in integer
paise so no lossy float arithmetic is used for audit values.

Text cells are escaped before writing so a value such as
``=HYPERLINK("...")`` can never become an executable spreadsheet formula.
After writing, the workbook is reopened and its sheets, counts, totals and
run metadata are checked against the database.

Excel does not permit ``/`` in worksheet names. The requested
``Duplicate/Near-match Review`` sheet is therefore stored as
``Duplicate Near-match Review``; every other required sheet keeps its name.
"""

from __future__ import annotations

import io
import json
import re
from datetime import datetime, timezone
from typing import Any

import db
import reconciliation

FORMAT_VERSION = "milaan-reconciliation-workbook/1"

REQUIRED_SHEETS = (
    "Summary",
    "Matched",
    "Purchase Missing from 2B",
    "2B Missing from Books",
    "Field Discrepancies",
    "Duplicate Near-match Review",
    "Rejected Rows",
    "ITC Evidence",
    "Audit Metadata",
)

DISCREPANCY_REASONS = frozenset({
    reconciliation.REASON_TAXABLE_VALUE_MISMATCH,
    reconciliation.REASON_COMPONENT_MISMATCH,
    reconciliation.REASON_CESS_MISMATCH,
    reconciliation.REASON_DATE_MISMATCH,
    reconciliation.REASON_DATE_PERIOD_INCONSISTENT,
    reconciliation.REASON_AMOUNT_DIFFERENCE_BEYOND_TOLERANCE,
    reconciliation.REASON_AMOUNT_WITHIN_TOLERANCE,
})

_FORMULA_PREFIXES = ("=", "+", "-", "@")
_CONTROL_PREFIXES = ("\t", "\r", "\n")
_ILLEGAL_CELL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_MAX_AUDIT_TEXT = 30_000

RESULT_COLUMNS = (
    ("run_id", "Run ID"),
    ("status", "Status"),
    ("method", "Matching method"),
    ("reason_codes", "Reason codes"),
    ("invoice_id", "Invoice ID"),
    ("supplier_name", "Supplier"),
    ("invoice_number", "Invoice number"),
    ("invoice_number_normalized", "Invoice number (normalized)"),
    ("supplier_gstin", "Supplier GSTIN"),
    ("tax_period", "Tax period"),
    ("invoice_date", "Invoice date"),
    ("invoice_taxable_rupees", "Invoice taxable (Rs)"),
    ("invoice_cgst_rupees", "Invoice CGST (Rs)"),
    ("invoice_sgst_rupees", "Invoice SGST (Rs)"),
    ("invoice_igst_rupees", "Invoice IGST (Rs)"),
    ("invoice_cess_rupees", "Invoice cess (Rs)"),
    ("invoice_total_gst_rupees", "Invoice total GST (Rs)"),
    ("twob_row_id", "2B row ID"),
    ("twob_row_number", "2B file row"),
    ("twob_invoice_number", "2B invoice number"),
    ("twob_invoice_number_normalized", "2B invoice number (normalized)"),
    ("twob_taxable_rupees", "2B taxable (Rs)"),
    ("twob_cgst_rupees", "2B CGST (Rs)"),
    ("twob_sgst_rupees", "2B SGST (Rs)"),
    ("twob_igst_rupees", "2B IGST (Rs)"),
    ("twob_cess_rupees", "2B cess (Rs)"),
    ("twob_total_gst_rupees", "2B total GST (Rs)"),
    ("amount_difference_paise", "Difference total GST (paise)"),
    ("itc_availability", "ITC availability"),
    ("itc_reason", "ITC reason"),
    ("reverse_charge", "Reverse charge"),
    ("document_type", "Document type"),
    ("irn", "IRN"),
    ("source_sheet", "2B sheet"),
    ("source_row", "2B sheet row"),
)

DISCREPANCY_COLUMNS = RESULT_COLUMNS + (
    ("fields_with_differences", "Fields with differences"),
    ("taxable_difference_paise", "Taxable difference (paise)"),
    ("cgst_difference_paise", "CGST difference (paise)"),
    ("sgst_difference_paise", "SGST difference (paise)"),
    ("igst_difference_paise", "IGST difference (paise)"),
    ("cess_difference_paise", "Cess difference (paise)"),
    ("total_gst_difference_paise", "Total GST difference (paise)"),
    ("date_difference", "Date difference"),
)

REVIEW_COLUMNS = RESULT_COLUMNS + (
    ("review_category", "Review category"),
)

REJECTED_COLUMNS = (
    ("source", "Source"),
    ("source_id", "Source ID"),
    ("batch_or_run", "Batch / run"),
    ("source_sheet", "Sheet"),
    ("source_row", "Source row"),
    ("status", "Status"),
    ("gstin", "GSTIN"),
    ("invoice_number", "Invoice number"),
    ("invoice_date", "Invoice date"),
    ("tax_period", "Tax period"),
    ("taxable_rupees", "Taxable (Rs)"),
    ("total_gst_rupees", "Total GST (Rs)"),
    ("reason_codes", "Reason codes"),
    ("issues", "Rejection reasons"),
    ("raw_source", "Raw source row"),
)

ITC_COLUMNS = (
    ("run_id", "Run ID"),
    ("status", "Match status"),
    ("invoice_number", "Invoice number"),
    ("supplier_gstin", "Supplier GSTIN"),
    ("tax_period", "Tax period"),
    ("invoice_total_gst_rupees", "Invoice total GST (Rs)"),
    ("twob_row_id", "2B row ID"),
    ("twob_total_gst_rupees", "2B total GST (Rs)"),
    ("itc_availability", "ITC availability"),
    ("itc_reason", "ITC reason"),
    ("reverse_charge", "Reverse charge"),
    ("document_type", "Document type"),
    ("irn", "IRN"),
    ("source_sheet", "2B sheet"),
    ("source_row", "2B sheet row"),
    ("evidence_note", "Evidence note"),
)


class WorkbookValidationError(ValueError):
    """Raised when a generated workbook fails its own reopen checks."""


def _looks_like_formula(text: str) -> bool:
    stripped = text.lstrip()
    return bool(stripped) and stripped[0] in _FORMULA_PREFIXES


def safe_cell(value: Any) -> Any:
    """Escape spreadsheet formula triggers and unreadable control chars."""
    if value is None or isinstance(value, (int, float, bool)):
        return value
    text = _ILLEGAL_CELL_CHARS.sub(" ", str(value))
    if (
        (text and text[0] in _CONTROL_PREFIXES)
        or _looks_like_formula(text)
    ):
        return "'" + text
    if len(text) > _MAX_AUDIT_TEXT:
        return text[:_MAX_AUDIT_TEXT] + "…[truncated]"
    return text


def rupees(paise: Any) -> float | None:
    if paise is None:
        return None
    return round(int(paise) / 100, 2)


def _codes(raw: Any) -> tuple[str, ...]:
    if isinstance(raw, (list, tuple)):
        return tuple(str(item) for item in raw)
    try:
        value = json.loads(raw or "[]")
    except (TypeError, ValueError):
        return (str(raw),) if raw else ()
    if not isinstance(value, list):
        return (str(value),)
    return tuple(str(item) for item in value)


def _issues(raw: Any) -> str:
    try:
        value = json.loads(raw or "[]")
    except (TypeError, ValueError):
        return str(raw or "")
    if not isinstance(value, list):
        return str(value)
    return "; ".join(
        f"[{item.get('code', '')}] {item.get('field', '')}: "
        f"{item.get('message', '')}"
        for item in value
        if isinstance(item, dict)
    )


def _load(run_id: int) -> dict:
    run = db.get_reconciliation_run(run_id)
    if run is None:
        raise ValueError(f"Reconciliation run {run_id} not found")
    batch = db.get_gstr2b_import(run["gstr2b_import_id"])
    if batch is None:
        raise ValueError(
            f"GSTR-2B import {run['gstr2b_import_id']} for run {run_id} not found"
        )
    with db.db() as conn:
        result_rows = conn.execute(
            """SELECT rr.*,
                      i.invoice_number AS books_invoice_number,
                      i.invoice_date AS books_invoice_date,
                      i.taxable_value AS books_taxable,
                      i.cgst AS books_cgst,
                      i.sgst AS books_sgst,
                      i.igst AS books_igst,
                      i.total_gst AS books_total_gst,
                      i.validation_status AS books_validation_status,
                      s.supplier_name,
                      s.gstin AS books_gstin,
                      gr.row_number AS twob_row_number,
                      gr.invoice_number AS twob_invoice_number,
                      gr.invoice_date AS twob_invoice_date,
                      gr.tax_period AS twob_tax_period,
                      gr.taxable_value AS twob_taxable,
                      gr.cgst AS twob_cgst,
                      gr.sgst AS twob_sgst,
                      gr.igst AS twob_igst,
                      gr.cess AS twob_cess,
                      gr.total_gst AS twob_total_gst,
                      gr.itc_availability AS twob_itc_availability,
                      gr.itc_reason AS twob_itc_reason,
                      gr.reverse_charge AS twob_reverse_charge,
                      gr.document_type AS twob_document_type,
                      gr.irn AS twob_irn,
                      gr.source_sheet AS twob_source_sheet,
                      gr.source_row AS twob_source_row
               FROM reconciliation_results rr
               LEFT JOIN invoices i ON i.id = rr.invoice_id
               LEFT JOIN suppliers s ON s.id = i.supplier_id
               LEFT JOIN gstr2b_rows gr ON gr.id = rr.gstr2b_row_id
               WHERE rr.run_id = ?
               ORDER BY COALESCE(rr.invoice_id, 0), rr.id""",
            (run_id,),
        ).fetchall()
        rejected_rows = conn.execute(
            """SELECT * FROM gstr2b_rows
               WHERE import_id = ? AND validation_status != 'valid'
               ORDER BY row_number""",
            (run["gstr2b_import_id"],),
        ).fetchall()
        totals = conn.execute(
            """SELECT
                   COALESCE(SUM(CASE WHEN invoice_id IS NOT NULL
                                     THEN invoice_total_gst_paise END), 0)
                       AS invoice_total,
                   COALESCE(SUM(CASE WHEN match_status = 'MATCHED'
                                     THEN invoice_total_gst_paise END), 0)
                       AS matched_total,
                   COALESCE(SUM(CASE WHEN match_status = 'MISSING_FROM_2B'
                                     THEN invoice_total_gst_paise END), 0)
                       AS missing_total,
                   COALESCE(SUM(CASE WHEN match_status = 'AMOUNT_MISMATCH'
                                     THEN invoice_total_gst_paise END), 0)
                       AS amount_mismatch_total,
                   COALESCE(SUM(CASE WHEN match_status = 'NEEDS_REVIEW'
                                     THEN invoice_total_gst_paise END), 0)
                       AS needs_review_total,
                   COALESCE(SUM(CASE WHEN match_status = 'NOT_ELIGIBLE'
                                     THEN invoice_total_gst_paise END), 0)
                       AS not_eligible_total,
                   COALESCE(SUM(CASE WHEN match_status = 'ITC_UNAVAILABLE'
                                     THEN invoice_total_gst_paise END), 0)
                       AS itc_unavailable_total,
                   COALESCE(SUM(CASE WHEN match_status = 'IN_2B_NOT_IN_BOOKS'
                                     THEN gstr2b_total_gst_paise END), 0)
                       AS in_2b_not_in_books_total
               FROM reconciliation_results WHERE run_id = ?""",
            (run_id,),
        ).fetchone()

    results = [dict(row) for row in result_rows]
    for item in results:
        item["reason_codes_list"] = _codes(item.get("reason_codes"))
    return {
        "run": run,
        "batch": batch,
        "results": results,
        "rejected_rows": [dict(row) for row in rejected_rows],
        "totals": dict(totals),
    }


def _field_differences(item: dict) -> tuple[list[str], dict[str, int | None], str]:
    fields: list[str] = []
    differences: dict[str, int | None] = {}
    date_parts: list[str] = []

    left_gstin = reconciliation.normalize_gstin(item.get("books_gstin"))
    right_gstin = reconciliation.normalize_gstin(item.get("supplier_gstin"))
    if left_gstin and right_gstin and left_gstin != right_gstin:
        fields.append("GSTIN")

    left_number = reconciliation.normalize_invoice_number(
        item.get("books_invoice_number")
    )
    right_number = reconciliation.normalize_invoice_number(
        item.get("twob_invoice_number")
    )
    if left_number and right_number and left_number != right_number:
        fields.append("Invoice number")

    left_period = reconciliation.normalize_tax_period(item.get("tax_period"))
    right_period = reconciliation.normalize_tax_period(
        item.get("twob_tax_period")
    )
    if left_period and right_period and left_period != right_period:
        fields.append("Tax period")

    left_date = item.get("books_invoice_date")
    right_date = item.get("twob_invoice_date")
    if left_date and right_date and left_date != right_date:
        fields.append("Invoice date")
        date_parts.append(f"{left_date} vs {right_date}")

    pairs = (
        ("Taxable value", "books_taxable", "twob_taxable",
         "taxable_difference_paise"),
        ("CGST", "books_cgst", "twob_cgst", "cgst_difference_paise"),
        ("SGST", "books_sgst", "twob_sgst", "sgst_difference_paise"),
        ("IGST", "books_igst", "twob_igst", "igst_difference_paise"),
        ("Total GST", "books_total_gst", "twob_total_gst",
         "total_gst_difference_paise"),
    )
    for label, left_key, right_key, difference_key in pairs:
        left = item.get(left_key)
        right = item.get(right_key)
        if left is not None and right is not None:
            difference = int(left) - int(right)
            differences[difference_key] = difference
            if difference != 0:
                fields.append(label)
        else:
            differences[difference_key] = None

    if item.get("books_cess") is not None and item.get("twob_cess") is not None:
        cess_difference = int(item["books_cess"]) - int(item["twob_cess"])
        differences["cess_difference_paise"] = cess_difference
        if cess_difference != 0:
            fields.append("Cess")
    else:
        # Purchase invoices do not currently store a cess amount. Do not
        # invent a zero and report a phantom mismatch against 2B cess.
        differences["cess_difference_paise"] = None

    return fields, differences, "; ".join(date_parts)


def _is_discrepancy(item: dict) -> bool:
    if not item.get("gstr2b_row_id") or item.get("invoice_id") is None:
        return False
    fields, _differences, _date = _field_differences(item)
    if fields:
        return True
    return bool(DISCREPANCY_REASONS.intersection(item["reason_codes_list"]))


def _result_display(item: dict) -> dict:
    fields, differences, date_difference = _field_differences(item)
    return {
        "run_id": item["run_id"],
        "status": item["match_status"],
        "method": item["match_method"],
        "reason_codes": "; ".join(item["reason_codes_list"]),
        "invoice_id": item.get("invoice_id"),
        "supplier_name": item.get("supplier_name") or "",
        "invoice_number": item.get("books_invoice_number")
        or item.get("invoice_number"),
        "invoice_number_normalized": reconciliation.normalize_invoice_number(
            item.get("books_invoice_number") or item.get("invoice_number")
        ),
        "supplier_gstin": item.get("supplier_gstin") or item.get("books_gstin"),
        "tax_period": item.get("tax_period"),
        "invoice_date": item.get("books_invoice_date"),
        "invoice_taxable_rupees": rupees(item.get("books_taxable")),
        "invoice_cgst_rupees": rupees(item.get("books_cgst")),
        "invoice_sgst_rupees": rupees(item.get("books_sgst")),
        "invoice_igst_rupees": rupees(item.get("books_igst")),
        "invoice_cess_rupees": rupees(item.get("books_cess") or 0),
        "invoice_total_gst_rupees": rupees(item.get("books_total_gst")),
        "twob_row_id": item.get("gstr2b_row_id"),
        "twob_row_number": item.get("twob_row_number"),
        "twob_invoice_number": item.get("twob_invoice_number"),
        "twob_invoice_number_normalized": reconciliation.normalize_invoice_number(
            item.get("twob_invoice_number")
        ),
        "twob_taxable_rupees": rupees(item.get("twob_taxable")),
        "twob_cgst_rupees": rupees(item.get("twob_cgst")),
        "twob_sgst_rupees": rupees(item.get("twob_sgst")),
        "twob_igst_rupees": rupees(item.get("twob_igst")),
        "twob_cess_rupees": rupees(item.get("twob_cess") or 0),
        "twob_total_gst_rupees": rupees(item.get("twob_total_gst")),
        "amount_difference_paise": item.get("amount_difference_paise"),
        "itc_availability": item.get("twob_itc_availability") or "",
        "itc_reason": item.get("twob_itc_reason") or "",
        "reverse_charge": item.get("twob_reverse_charge") or "",
        "document_type": item.get("twob_document_type") or "",
        "irn": item.get("twob_irn") or "",
        "source_sheet": item.get("twob_source_sheet") or "",
        "source_row": item.get("twob_source_row"),
        "fields_with_differences": ", ".join(fields),
        "taxable_difference_paise": differences.get("taxable_difference_paise"),
        "cgst_difference_paise": differences.get("cgst_difference_paise"),
        "sgst_difference_paise": differences.get("sgst_difference_paise"),
        "igst_difference_paise": differences.get("igst_difference_paise"),
        "cess_difference_paise": differences.get("cess_difference_paise"),
        "total_gst_difference_paise": differences.get(
            "total_gst_difference_paise"
        ),
        "date_difference": date_difference,
        "review_category": _review_category(item["reason_codes_list"]),
    }


def _review_category(codes: tuple[str, ...]) -> str:
    if any(code.startswith("DUPLICATE") for code in codes):
        return "Duplicate / ambiguous"
    if reconciliation.REASON_NEAR_INVOICE_NUMBER in codes:
        return "Near invoice number"
    if reconciliation.REASON_AMBIGUOUS_CANDIDATES in codes:
        return "Ambiguous candidate"
    if reconciliation.REASON_DATE_PERIOD_INCONSISTENT in codes:
        return "Date / period review"
    if reconciliation.REASON_CANDIDATE_NOT_EXACT_INVOICE_NUMBER in codes:
        return "Amount-only candidate"
    return "Other review"


def _itc_note(item: dict) -> str:
    availability = (item.get("twob_itc_availability") or "").upper()
    reverse = (item.get("twob_reverse_charge") or "").upper()
    if availability in {"N", "NO", "UNAVAILABLE"}:
        return (
            "2B marks ITC unavailable or restricted. Do not claim without "
            "trader/CA review."
        )
    if reverse in {"Y", "YES"}:
        return (
            "2B marks reverse charge. The recipient may be liable; review "
            "with the trader/CA."
        )
    return (
        "Invoice appears in the imported 2B batch. Presence in 2B does not "
        "by itself establish ITC eligibility."
    )


def _rejected_display(data: dict) -> list[dict]:
    rows: list[dict] = []
    for item in data["rejected_rows"]:
        rows.append({
            "source": "GSTR-2B row",
            "source_id": item.get("id"),
            "batch_or_run": data["run"]["gstr2b_import_id"],
            "source_sheet": item.get("source_sheet") or "",
            "source_row": item.get("source_row") or item.get("row_number"),
            "status": item.get("validation_status"),
            "gstin": item.get("supplier_gstin") or "",
            "invoice_number": item.get("invoice_number") or "",
            "invoice_date": item.get("invoice_date") or "",
            "tax_period": item.get("tax_period") or "",
            "taxable_rupees": rupees(item.get("taxable_value")),
            "total_gst_rupees": rupees(item.get("total_gst")),
            "reason_codes": "; ".join(
                issue.get("code", "")
                for issue in _json_list(item.get("errors"))
                if isinstance(issue, dict)
            ),
            "issues": _issues(item.get("errors")),
            "raw_source": item.get("raw_data") or "",
        })
    for item in data["results"]:
        if item["match_status"] != reconciliation.STATUS_NOT_ELIGIBLE:
            continue
        rows.append({
            "source": "Purchase invoice",
            "source_id": item.get("invoice_id"),
            "batch_or_run": data["run"]["id"],
            "source_sheet": "",
            "source_row": "",
            "status": item.get("books_validation_status")
            or reconciliation.STATUS_NOT_ELIGIBLE,
            "gstin": item.get("supplier_gstin") or "",
            "invoice_number": item.get("books_invoice_number")
            or item.get("invoice_number") or "",
            "invoice_date": item.get("books_invoice_date") or "",
            "tax_period": item.get("tax_period") or "",
            "taxable_rupees": rupees(item.get("books_taxable")),
            "total_gst_rupees": rupees(item.get("books_total_gst")),
            "reason_codes": "; ".join(item["reason_codes_list"]),
            "issues": "; ".join(item["reason_codes_list"]),
            "raw_source": "",
        })
    return rows


def _json_list(raw: Any) -> list:
    try:
        value = json.loads(raw or "[]")
    except (TypeError, ValueError):
        return []
    return value if isinstance(value, list) else []


def _append_table(worksheet, columns, rows) -> None:
    worksheet.append([safe_cell(label) for _key, label in columns])
    for row in rows:
        worksheet.append([safe_cell(row.get(key)) for key, _label in columns])


def _style_table(worksheet, widths: dict[int, int] | None = None) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F3864")
    for cell in worksheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for index in range(1, worksheet.max_column + 1):
        letter = get_column_letter(index)
        width = (widths or {}).get(index, 20)
        worksheet.column_dimensions[letter].width = width


def _build_summary(worksheet, data: dict, generated_at: str) -> None:
    from openpyxl.styles import Font

    run = data["run"]
    batch = data["batch"]
    totals = data["totals"]
    title = "MILAAN GSTR-2B reconciliation workbook"
    worksheet.append([title])
    worksheet["A1"].font = Font(bold=True, size=14)
    worksheet.append([])
    rows = (
        ("Run ID", run["id"]),
        ("Run created at", run["created_at"]),
        ("GSTR-2B batch", run["gstr2b_import_id"]),
        ("Source file", batch["file_name"]),
        ("Source SHA-256", batch["file_hash"]),
        ("Imported at", batch["imported_at"]),
        ("Source format", batch["source_format"]),
        ("Parser", batch["parser_kind"]),
        ("Amount tolerance (paise)", run["amount_tolerance_paise"]),
        ("Invoices compared", run["invoice_count"]),
        ("Matched", run["matched_count"]),
        ("Amount mismatch", run["amount_mismatch_count"]),
        ("Purchase missing from 2B", run["missing_count"]),
        ("2B rows missing from books", run["in_2b_not_in_books_count"]),
        ("Needs review", run["needs_review_count"]),
        ("Not eligible", run["not_eligible_count"]),
        ("ITC unavailable", run["itc_unavailable_count"]),
        ("Invoice total GST (paise)", totals["invoice_total"]),
        ("Matched total GST (paise)", totals["matched_total"]),
        ("Missing total GST (paise)", totals["missing_total"]),
        ("Amount mismatch total GST (paise)",
         totals["amount_mismatch_total"]),
        ("Needs review total GST (paise)", totals["needs_review_total"]),
        ("Not eligible total GST (paise)", totals["not_eligible_total"]),
        ("ITC unavailable total GST (paise)",
         totals["itc_unavailable_total"]),
        ("2B-only total GST (paise)", totals["in_2b_not_in_books_total"]),
        ("Report generated at", generated_at),
    )
    for label, value in rows:
        worksheet.append([safe_cell(label), safe_cell(value)])

    worksheet.append([])
    worksheet.append(["Purpose", "Deterministic two-way reconciliation evidence."])
    worksheet.append([
        "Monetary unit",
        "Rupee columns use INR with two decimals; difference columns are "
        "integer paise (100 paise = 1 rupee).",
    ])
    worksheet.append([
        "Matching policy",
        "Exact GSTIN + normalized invoice number + tax period first. Near "
        "matches are review-only. No cross-GSTIN or cross-period auto-match.",
    ])
    worksheet.append([
        "ITC caution",
        "Presence in GSTR-2B does not by itself establish ITC eligibility. "
        "Uncertain cases require trader/CA review.",
    ])
    worksheet.append([
        "Recovery caution",
        "A sent message or supplier promise is not recovery evidence. "
        "Recovery requires later GSTR-2B or other supported evidence.",
    ])
    worksheet.append([])
    worksheet.append(["Workbook sheets"])
    worksheet.append(["Sheet", "Data rows"])
    for sheet_name in REQUIRED_SHEETS[1:-1]:
        worksheet.append([safe_cell(sheet_name), safe_cell(
            _expected_sheet_counts(data).get(sheet_name, 0)
        )])
    worksheet.column_dimensions["A"].width = 34
    worksheet.column_dimensions["B"].width = 90


def _expected_sheet_counts(data: dict) -> dict[str, int]:
    run = data["run"]
    discrepancies = [
        item for item in data["results"] if _is_discrepancy(item)
    ]
    return {
        "Matched": run["matched_count"],
        "Purchase Missing from 2B": run["missing_count"],
        "2B Missing from Books": run["in_2b_not_in_books_count"],
        "Field Discrepancies": len(discrepancies),
        "Duplicate Near-match Review": run["needs_review_count"],
        "Rejected Rows": len(data["rejected_rows"]) + run["not_eligible_count"],
        "ITC Evidence": sum(
            1 for item in data["results"] if item.get("gstr2b_row_id")
        ),
    }


def _build_audit(worksheet, data: dict, generated_at: str) -> None:
    run = data["run"]
    batch = data["batch"]
    counts = _expected_sheet_counts(data)
    worksheet.append(["Key", "Value"])
    rows = (
        ("Format version", FORMAT_VERSION),
        ("Generated at", generated_at),
        ("Run ID", run["id"]),
        ("Run created at", run["created_at"]),
        ("GSTR-2B batch", run["gstr2b_import_id"]),
        ("Source file", batch["file_name"]),
        ("Source SHA-256", batch["file_hash"]),
        ("Source size (bytes)", batch["file_size"]),
        ("Source format", batch["source_format"]),
        ("Parser kind", batch["parser_kind"]),
        ("Imported at", batch["imported_at"]),
        ("Sheets used", batch.get("sheets_used")),
        ("Skipped sheets", batch.get("skipped_sheets")),
        ("Unsupported fields", batch.get("unsupported_fields")),
        ("Amount tolerance (paise)", run["amount_tolerance_paise"]),
        ("Required sheets", json.dumps(list(REQUIRED_SHEETS))),
        ("Sheet row counts", json.dumps(counts, sort_keys=True)),
        ("Report totals", json.dumps(data["totals"], sort_keys=True)),
        ("Selection policy", (
            "Exact key match first; near/amount-only candidates review-only; "
            "no cross-GSTIN or cross-period matching; duplicate keys never "
            "auto-selected."
        )),
        ("Formula protection", (
            "Untrusted text beginning with =, +, -, @, tab or carriage return "
            "is prefixed with an apostrophe."
        )),
    )
    for label, value in rows:
        worksheet.append([safe_cell(label), safe_cell(value)])
    worksheet.column_dimensions["A"].width = 28
    worksheet.column_dimensions["B"].width = 110


def build_reconciliation_workbook(run_id: int, *, validate: bool = True) -> bytes:
    """Build and self-validate the workbook for one reconciliation run."""
    from openpyxl import Workbook

    if not isinstance(run_id, int) or isinstance(run_id, bool):
        raise ValueError("run_id must be an integer")
    data = _load(run_id)
    generated_at = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")

    workbook = Workbook()
    workbook.remove(workbook.active)

    summary = workbook.create_sheet("Summary")
    _build_summary(summary, data, generated_at)

    displays = [_result_display(item) for item in data["results"]]
    matched = [row for row in displays if row["status"] == "MATCHED"]
    missing = [
        row for row in displays if row["status"] == "MISSING_FROM_2B"
    ]
    twob_only = [
        row for row in displays if row["status"] == "IN_2B_NOT_IN_BOOKS"
    ]
    discrepancies = [
        _result_display(item)
        for item in data["results"] if _is_discrepancy(item)
    ]
    review = [
        row for row in displays if row["status"] == "NEEDS_REVIEW"
    ]
    itc_evidence = []
    for item in data["results"]:
        if not item.get("gstr2b_row_id"):
            continue
        row = _result_display(item)
        row["evidence_note"] = _itc_note(item)
        itc_evidence.append(row)

    sheets = (
        ("Matched", RESULT_COLUMNS, matched, {}),
        ("Purchase Missing from 2B", RESULT_COLUMNS, missing, {}),
        ("2B Missing from Books", RESULT_COLUMNS, twob_only, {}),
        ("Field Discrepancies", DISCREPANCY_COLUMNS, discrepancies, {}),
        ("Duplicate Near-match Review", REVIEW_COLUMNS, review, {}),
        ("Rejected Rows", REJECTED_COLUMNS, _rejected_display(data), {}),
        ("ITC Evidence", ITC_COLUMNS, itc_evidence, {}),
    )
    for name, columns, rows, widths in sheets:
        worksheet = workbook.create_sheet(name)
        _append_table(worksheet, columns, rows)
        _style_table(worksheet, widths)

    audit = workbook.create_sheet("Audit Metadata")
    _build_audit(audit, data, generated_at)

    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    content = buffer.getvalue()

    if validate:
        problems = validate_workbook(content, run_id)
        if problems:
            raise WorkbookValidationError(
                "Generated workbook failed validation: " + "; ".join(problems)
            )
    return content


def _sheet_rows(worksheet) -> int:
    return max(worksheet.max_row - 1, 0)


def _read_key_values(worksheet, allowed: set[str]) -> dict[str, Any]:
    values: dict[str, Any] = {}
    for row in worksheet.iter_rows(values_only=True):
        if not row or len(row) < 2:
            continue
        key = row[0]
        if isinstance(key, str) and key in allowed:
            values[key] = row[1]
    return values


def _expected_totals(run_id: int) -> dict[str, int]:
    with db.db() as conn:
        row = conn.execute(
            """SELECT
                   COALESCE(SUM(CASE WHEN invoice_id IS NOT NULL
                                     THEN invoice_total_gst_paise END), 0)
                       AS invoice_total,
                   COALESCE(SUM(CASE WHEN match_status = 'MATCHED'
                                     THEN invoice_total_gst_paise END), 0)
                       AS matched_total,
                   COALESCE(SUM(CASE WHEN match_status = 'MISSING_FROM_2B'
                                     THEN invoice_total_gst_paise END), 0)
                       AS missing_total,
                   COALESCE(SUM(CASE WHEN match_status = 'AMOUNT_MISMATCH'
                                     THEN invoice_total_gst_paise END), 0)
                       AS amount_mismatch_total,
                   COALESCE(SUM(CASE WHEN match_status = 'NEEDS_REVIEW'
                                     THEN invoice_total_gst_paise END), 0)
                       AS needs_review_total,
                   COALESCE(SUM(CASE WHEN match_status = 'NOT_ELIGIBLE'
                                     THEN invoice_total_gst_paise END), 0)
                       AS not_eligible_total,
                   COALESCE(SUM(CASE WHEN match_status = 'ITC_UNAVAILABLE'
                                     THEN invoice_total_gst_paise END), 0)
                       AS itc_unavailable_total,
                   COALESCE(SUM(CASE WHEN match_status = 'IN_2B_NOT_IN_BOOKS'
                                     THEN gstr2b_total_gst_paise END), 0)
                       AS in_2b_not_in_books_total
               FROM reconciliation_results WHERE run_id = ?""",
            (run_id,),
        ).fetchone()
    return {key: int(row[key]) for key in row.keys()}


def validate_workbook(content: bytes, run_id: int) -> list[str]:
    """Reopen a generated workbook and compare it with stored run evidence."""
    from openpyxl import load_workbook

    run = db.get_reconciliation_run(run_id)
    if run is None:
        return [f"run {run_id} not found"]
    try:
        # Do not use data_only=True here: cached values would hide a formula
        # injected into a cell, defeating the formula-safety check below.
        workbook = load_workbook(io.BytesIO(content), data_only=False)
    except Exception as exc:
        return [f"workbook cannot be reopened: {type(exc).__name__}"]

    problems: list[str] = []
    try:
        missing_sheets = [
            name for name in REQUIRED_SHEETS if name not in workbook.sheetnames
        ]
        if missing_sheets:
            problems.append("missing sheets: " + ", ".join(missing_sheets))
            return problems

        expected_counts = _expected_sheet_counts(_load(run_id))
        for sheet_name, expected in expected_counts.items():
            actual = _sheet_rows(workbook[sheet_name])
            if actual != expected:
                problems.append(
                    f"{sheet_name} rows {actual} != expected {expected}"
                )

        summary = _read_key_values(workbook["Summary"], {
            "Run ID", "GSTR-2B batch", "Source file", "Source SHA-256",
            "Amount tolerance (paise)", "Invoices compared", "Matched",
            "Amount mismatch", "Purchase missing from 2B",
            "2B rows missing from books", "Needs review", "Not eligible",
            "ITC unavailable", "Invoice total GST (paise)",
            "Matched total GST (paise)", "Missing total GST (paise)",
            "Amount mismatch total GST (paise)",
            "Needs review total GST (paise)",
            "Not eligible total GST (paise)",
            "ITC unavailable total GST (paise)",
            "2B-only total GST (paise)",
        })
        expected_summary = {
            "Run ID": run["id"],
            "GSTR-2B batch": run["gstr2b_import_id"],
            "Amount tolerance (paise)": run["amount_tolerance_paise"],
            "Invoices compared": run["invoice_count"],
            "Matched": run["matched_count"],
            "Amount mismatch": run["amount_mismatch_count"],
            "Purchase missing from 2B": run["missing_count"],
            "2B rows missing from books": run["in_2b_not_in_books_count"],
            "Needs review": run["needs_review_count"],
            "Not eligible": run["not_eligible_count"],
            "ITC unavailable": run["itc_unavailable_count"],
        }
        for key, expected in expected_summary.items():
            if summary.get(key) != expected:
                problems.append(
                    f"summary {key} {summary.get(key)!r} != {expected!r}"
                )

        expected_totals = _expected_totals(run_id)
        total_labels = {
            "invoice_total": "Invoice total GST (paise)",
            "matched_total": "Matched total GST (paise)",
            "missing_total": "Missing total GST (paise)",
            "amount_mismatch_total": "Amount mismatch total GST (paise)",
            "needs_review_total": "Needs review total GST (paise)",
            "not_eligible_total": "Not eligible total GST (paise)",
            "itc_unavailable_total": "ITC unavailable total GST (paise)",
            "in_2b_not_in_books_total": "2B-only total GST (paise)",
        }
        for key, label in total_labels.items():
            if summary.get(label) != expected_totals[key]:
                problems.append(
                    f"summary {label} {summary.get(label)!r} != "
                    f"{expected_totals[key]!r}"
                )

        audit = _read_key_values(
            workbook["Audit Metadata"],
            {"Format version", "Report totals", "Sheet row counts", "Run ID"},
        )
        if audit.get("Format version") != FORMAT_VERSION:
            problems.append("audit format version mismatch")
        if audit.get("Run ID") != run_id:
            problems.append("audit run id mismatch")
        try:
            audit_totals = json.loads(audit.get("Report totals") or "{}")
        except (TypeError, ValueError):
            audit_totals = {}
            problems.append("audit totals are not valid JSON")
        for key, expected in expected_totals.items():
            if audit_totals.get(key) != expected:
                problems.append(
                    f"audit total {key} {audit_totals.get(key)!r} != "
                    f"{expected!r}"
                )
        try:
            audit_counts = json.loads(audit.get("Sheet row counts") or "{}")
        except (TypeError, ValueError):
            audit_counts = {}
            problems.append("audit sheet counts are not valid JSON")
        for key, expected in expected_counts.items():
            if audit_counts.get(key) != expected:
                problems.append(
                    f"audit count {key} {audit_counts.get(key)!r} != "
                    f"{expected!r}"
                )

        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows(values_only=True):
                for value in row:
                    if not isinstance(value, str) or not value:
                        continue
                    if value[0] in _CONTROL_PREFIXES or _looks_like_formula(
                        value
                    ):
                        problems.append(
                            f"unescaped formula-like value in "
                            f"{worksheet.title}"
                        )
    finally:
        workbook.close()

    return problems
